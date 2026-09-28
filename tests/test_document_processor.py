"""document_processor.py is pure, local, offline logic (text extraction + chunking)
— no network calls, no Pinecone, no LLM — so it's fully testable without any paid
API. Fixture files are generated on the fly rather than checked in as binaries.
"""

import pymupdf
import pytest
from docx import Document

from app.services import document_processor


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from a real PDF page.")
    doc.save(path)
    doc.close()
    return str(path)


@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("First paragraph of the document.")
    doc.add_paragraph("Second paragraph, with more detail.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    doc.save(path)
    return str(path)


@pytest.fixture
def sample_eml(tmp_path):
    path = tmp_path / "sample.eml"
    path.write_text(
        "Subject: Policy Update\n"
        "From: sender@example.com\n"
        "To: recipient@example.com\n"
        "Date: Mon, 1 Jan 2024 00:00:00 +0000\n"
        "Content-Type: text/plain\n"
        "\n"
        "This is the plain text body of the email.\n"
    )
    return str(path)


def test_extract_text_from_pdf(sample_pdf):
    text = document_processor.extract_text_from_pdf(sample_pdf)
    assert "Hello from a real PDF page." in text


def test_extract_text_from_docx(sample_docx):
    text = document_processor.extract_text_from_docx(sample_docx)
    assert "First paragraph of the document." in text
    assert "Second paragraph, with more detail." in text
    assert "Name | Value" in text


def test_extract_text_from_email(sample_eml):
    text = document_processor.extract_text_from_email(sample_eml)
    assert "Subject: Policy Update" in text
    assert "From: sender@example.com" in text
    assert "This is the plain text body of the email." in text


def test_extract_text_from_document_dispatches_by_extension(sample_pdf, sample_docx, sample_eml):
    assert "Hello from a real PDF page." in document_processor.extract_text_from_document(sample_pdf)
    assert "First paragraph" in document_processor.extract_text_from_document(sample_docx)
    assert "Subject: Policy Update" in document_processor.extract_text_from_document(sample_eml)


def test_extract_text_from_document_unsupported_extension_raises(tmp_path):
    path = tmp_path / "sample.xyz"
    path.write_text("irrelevant")
    with pytest.raises(ValueError, match="Unsupported file format"):
        document_processor.extract_text_from_document(str(path))


def test_chunk_text_splits_on_size():
    text = "\n".join(f"line {i}" for i in range(50))
    chunks = document_processor.chunk_text(text, chunk_size=50, overlap=10)
    assert len(chunks) > 1
    for chunk in chunks:
        # Each chunk may slightly exceed chunk_size by up to one line's worth
        # (a line is never split mid-way), but shouldn't balloon unbounded.
        assert len(chunk) < 100


def test_chunk_text_consecutive_chunks_actually_overlap():
    # Long enough that it's guaranteed to split into at least two chunks.
    text = "\n".join(f"This is line number {i} with some real content." for i in range(30))
    chunks = document_processor.chunk_text(text, chunk_size=200, overlap=40)
    assert len(chunks) >= 2

    # The whole point of `overlap` is that consecutive chunks share trailing/leading
    # content, so a sentence split at the boundary isn't lost from both sides.
    first, second = chunks[0], chunks[1]
    overlap_region = first[-40:].strip()
    assert (
        overlap_region.split()[-1] in second
    ), "expected the end of the first chunk to reappear at the start of the next"


def test_chunk_text_empty_input_returns_no_chunks():
    assert document_processor.chunk_text("") == []


def test_chunk_text_short_text_is_a_single_chunk():
    assert document_processor.chunk_text("short text", chunk_size=500) == ["short text"]
